"""Adapt the LV report package for MV without rewriting its layout or branding.

All changing values use the existing shared Variables API placeholders. Rebuild
after updating the LV template with the bundled Python runtime.
"""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'Template files' / 'ReportMacroTemplate_LV.docx'
TARGET = ROOT / 'Template files' / 'ReportMacroTemplate_MV.docx'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
OFFICE_REL_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
WP_NS = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'

REPLACEMENTS = {
    'Cable tray report': 'MV cable tray report',
    'This report records the routed cable schedule, explains the weight and installation-width calculations, and checks cable-plus-tray loading against the assigned manufacturer curve. Assessment results are collected before the final cable laying concept.':
        'This report records the medium voltage cable schedule, explains the weight and trefoil installation-width calculations, and checks cable-plus-tray loading against the assigned manufacturer curve. Enabled trefoil clamp mass is included in the carried load. Assessment results are collected before the final cable laying concept.',
    'The following schedule lists the cables routed through this tray, their diameters and their unit masses. These values form the input to the weight calculation. The cable laying concept on the final page shows their bundle arrangement.':
        'The following schedule lists the cables routed through this MV tray, their diameters and their unit masses. Each individual phase cable contributes once to the weight calculation. The cable laying concept on the final page shows the configured trefoil groups and any separately routed cables.',
    'Cables on tray weight calculations': 'Cables and trefoil clamps weight calculations',
    'The cable mass per metre is the sum of the unit masses of all the routed cables in the schedule. Each cable contributes once; the selected external grounding cable is added once when enabled.':
        'The carried cable mass per metre is the sum of the unit masses of all routed cables, plus the selected external grounding cable once when enabled. Enabled trefoil clamps contribute their equivalent mass per metre once. For route length L and clamp spacing s_c in metres, each clamped group uses n_c = ceil(L / s_c) + 1 clamps, including both ends. Total clamp mass is the sum of n_c multiplied by the selected clamp unit mass for each group; dividing by L gives the clamp mass per metre added as the final term in the sum below.',
    'Cables mass per metre: {{CABLEWEIGHTSUM}}':
        'Cables and enabled clamps mass per metre: {{CABLEWEIGHTSUM}}',
    'The total cable mass on the reported route is the cable mass per metre multiplied by the route length. Each routed cable is assumed to occupy this full length.':
        'Component summary: {{CABLESWEIGHTPERMETER}}. The combined cable and enabled clamp mass per metre is multiplied by the route length to obtain the carried mass on the route. Each routed cable is assumed to occupy this full length; clamp mass is already included and is not added again.',
    'Cables total mass: {{CABLESWEIGHTCALCULATIONS}}':
        'Cables and enabled clamps total mass: {{CABLESWEIGHTCALCULATIONS}}',
    'The combined route inventory adds the tray, support and cable masses. It is reported per metre and as a total for the route.':
        'The combined route inventory adds the tray, support and cable masses, including enabled trefoil clamps once in the cable component. It is reported per metre and as a total for the route.',
    'The uniformly distributed load used for the manufacturer curve consists of cable mass plus the own mass of the tray. The equivalent support mass used in the route inventory above is not included in this load.':
        'The uniformly distributed load used for the manufacturer curve consists of cable mass, including enabled trefoil clamps, plus the own mass of the tray. The equivalent support mass used in the route inventory above is not included in this load.',
    'Cable-plus-tray mass per metre: {{SPANMASS}}. The selected support spacing is {{CALCULATEDSPAN}} m.':
        'Cable-plus-tray mass per metre including enabled clamps: {{SPANMASS}}. The selected support spacing is {{CALCULATEDSPAN}} m.',
    'The combined cable and tray mass per metre is multiplied by gravitational acceleration g = 9.80665 m/s². Dividing by 1000 converts the resulting force per metre from N/m to kN/m.':
        'The combined cable, enabled clamp and tray mass per metre is multiplied by gravitational acceleration g = 9.80665 m/s². Dividing by 1000 converts the resulting force per metre from N/m to kN/m.',
    'Space occupied by cables': 'Space occupied by MV cable groups',
    'The overall tray side-rail height is {{TRAYHEIGHT}}. The useful height is found by subtracting the rung height from the overall height.':
        'Tray board height is {{TRAYHEIGHT}}, but the C-profiles occupy a part of the volume space. So, the useful height for the tray is {{USEFULTRAY}}. Medium voltage cables are laid and grouped in a triangle ("trefoil") formation, forming each a 3-phase system.',
    'For a single row, occupied width is the sum of cable diameters plus the spacing between adjacent cables. For bundles arranged in multiple rows, the calculation uses the bundle footprints from the actual Tray details layout; it does not add the diameters of every cable as though all cables were laid in one row.':
        'MV phase cables may be grouped in triangular trefoil formation, with three phase cables per group. Occupied width uses the actual group footprints from Tray details, including cable diameters, formation, rotation and the selected intergroup spacing. Where trefoil clamps are enabled, the selected clamp envelope is reserved in the footprint. Separately routed cables remain part of the layout. The three phase diameters are not simply added as though all cables occupied one row.',
    'The configured spacing between cables is {{CABLESPACING}}. The project option to count bundle spacing as free width is {{BUNDLESPACINGFREE}}. Cable widths and the spacing counted as occupied are combined as follows.':
        'The project cable-spacing setting is {{CABLESPACING}}. The installation requirement above specifies a minimum trefoil separation of 2d; the geometric calculation below reports the selected Tray details layout. The project option to count bundle spacing as free width is {{BUNDLESPACINGFREE}}. The formula adds the total group footprint width and the spacing counted as occupied. This free-width calculation does not verify the required separation or parallel cable lengths.',
    'This conclusion covers the specified distributed tray load and manufacturer curve. It does not establish the capacity of support brackets, fixings or anchors, or verify concentrated loads.':
        'This conclusion covers the specified distributed tray load and manufacturer curve. It does not establish the capacity of support brackets, fixings or anchors, or verify concentrated loads. It is not a short-circuit restraint or cable ampacity assessment; those checks require separate electrical and clamp design inputs.',
    'Picture 3  Cable distribution and bundle arrangement along the selected tray':
        'Picture 5  MV cable distribution trefoil groups and enabled clamps along the selected tray',
}


def insert_trefoil_details(document, source, relationships):
    """Retain the two original MV figures and the requested installation text."""
    anchors = document.xpath('//w:p[w:r/w:t="Useful tray height: {{USEFULTRAY}}"]', namespaces=NS)
    if len(anchors) != 1:
        raise ValueError('Expected one useful-height formula insertion point')
    draft = Document()
    for filename, caption in [
        ('mv-trefoil-formation.jpeg', 'Picture 3  Trefoil cables formation type overview'),
        ('mv-trefoil-spacing.jpeg', 'Picture 4  Minimum distance of cable trefoil'),
    ]:
        if 'spacing' in filename:
            text = draft.add_paragraph(
                'The minimum distance of cable bundle/trefoil is 2x outer cable diameter (2d). '
                'Between parallel laid power cables minimum distances must be complied with along '
                'the entire laying distance (except for building-pass-through / penetrations). '
                'Power cables laid in parallel shall have the same cable lengths. A maximum cable '
                'length difference of 3% can be accepted if properly considered in cable dimensioning '
                '(de-rating). No free space is considered.'
            )
            text.paragraph_format.keep_with_next = True
        image = draft.add_paragraph()
        image.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image.paragraph_format.keep_with_next = True
        image.add_run().add_picture(str(ROOT / 'Template files/report-assets' / filename), width=Mm(172))
        draft.add_paragraph(caption, 'Caption')

    ids = {rel.get('Id') for rel in relationships}
    drawing_ids = []
    for name in source.namelist():
        if name.startswith('word/') and name.endswith('.xml'):
            drawing_ids.extend(int(node.get('id')) for node in etree.fromstring(source.read(name)).iter(f'{{{WP_NS}}}docPr'))
    next_drawing_id = max(drawing_ids, default=0) + 1
    added = {}
    anchor = anchors[0]
    for paragraph in draft.paragraphs:
        copied = deepcopy(paragraph._p)
        for blip in copied.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            old_id = blip.get(f'{{{OFFICE_REL_NS}}}embed')
            image_part = draft.part.related_parts[old_id]
            filename = ['mv-trefoil-formation.jpeg', 'mv-trefoil-spacing.jpeg'][len(added)]
            new_id = f'rIdMvTrefoil{len(added) + 1}'
            if new_id in ids:
                raise ValueError(f'Duplicate figure relationship {new_id}')
            ids.add(new_id)
            rel = etree.SubElement(relationships, f'{{{REL_NS}}}Relationship')
            rel.set('Id', new_id)
            rel.set('Type', f'{OFFICE_REL_NS}/image')
            rel.set('Target', f'media/{filename}')
            blip.set(f'{{{OFFICE_REL_NS}}}embed', new_id)
            added[f'word/media/{filename}'] = image_part.blob
        for drawing in copied.iter(f'{{{WP_NS}}}docPr'):
            drawing.set('id', str(next_drawing_id))
            drawing.set('descr', 'Trefoil phase formation' if len(added) == 1 else 'Minimum 2d separation between trefoil groups')
            next_drawing_id += 1
        anchor.addnext(copied)
        anchor = copied
    return added


def build():
    with ZipFile(SOURCE) as source:
        document = etree.fromstring(source.read('word/document.xml'))
        counts = {text: 0 for text in REPLACEMENTS}
        for paragraph in document.xpath('//w:p', namespaces=NS):
            text_nodes = paragraph.xpath('.//w:t', namespaces=NS)
            original = ''.join(node.text or '' for node in text_nodes)
            if original not in REPLACEMENTS:
                continue
            # Keep paragraph/run formatting and all other package parts intact.
            text_nodes[0].text = REPLACEMENTS[original]
            for node in text_nodes[1:]:
                node.text = ''
            counts[original] += 1
        invalid = {text: count for text, count in counts.items() if count != 1}
        if invalid:
            raise ValueError(f'LV source changed; review the MV adaptation slots: {invalid}')

        relationships = etree.fromstring(source.read('word/_rels/document.xml.rels'))
        added = insert_trefoil_details(document, source, relationships)
        types = etree.fromstring(source.read('[Content_Types].xml'))
        content_ns = 'http://schemas.openxmlformats.org/package/2006/content-types'
        if not any(node.get('Extension') == 'jpeg' for node in types):
            image_type = etree.SubElement(types, f'{{{content_ns}}}Default')
            image_type.set('Extension', 'jpeg')
            image_type.set('ContentType', 'image/jpeg')

        core = etree.fromstring(source.read('docProps/core.xml'))
        for local, value in [
            ('title', 'MV cable tray report'),
            ('subject', 'MV tray cable and clamp weight load and free-width calculations'),
        ]:
            node = core.find(f'{{http://purl.org/dc/elements/1.1/}}{local}')
            if node is None:
                raise ValueError(f'Missing core property {local}')
            node.text = value

        updated = {
            'word/document.xml': etree.tostring(document, xml_declaration=True, encoding='UTF-8', standalone=True),
            'docProps/core.xml': etree.tostring(core, xml_declaration=True, encoding='UTF-8', standalone=True),
            'word/_rels/document.xml.rels': etree.tostring(relationships, xml_declaration=True, encoding='UTF-8', standalone=True),
            '[Content_Types].xml': etree.tostring(types, xml_declaration=True, encoding='UTF-8', standalone=True),
        }
        with ZipFile(TARGET, 'w') as target:
            for info in source.infolist():
                target.writestr(deepcopy(info), updated.get(info.filename, source.read(info.filename)))
            for name, data in added.items():
                target.writestr(name, data)
    print(TARGET)


if __name__ == '__main__':
    build()
