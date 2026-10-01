"""Build the LV template with the reference report's explanatory calculation flow."""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / 'Template files' / 'ReportMacroTemplate_LV.docx'
d = Document(TARGET)
for child in list(d._element.body):
    if child.tag != qn('w:sectPr'): d._element.body.remove(child)
s = d.sections[0]
s.page_width, s.page_height = Mm(210), Mm(297)
s.top_margin, s.bottom_margin = Mm(28), Mm(25)
s.left_margin, s.right_margin = Mm(19), Mm(19)
s.header_distance, s.footer_distance = Mm(8), Mm(5)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    if name not in d.styles: d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style = d.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor.from_string('000000')
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.1
normal = d.styles['Normal']; normal.font.size = Pt(11)
normal.paragraph_format.widow_control = True
for name, size in [('Title', 44), ('Subtitle', 22), ('Heading 1', 19), ('Heading 2', 15), ('Caption', 9)]:
    style = d.styles[name]; style.font.size = Pt(size)
    if name.startswith('Heading'):
        style.font.bold = False
        style.font.color.rgb = RGBColor.from_string('075374')
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.keep_with_next = True
    pp = style._element.find(qn('w:pPr'))
    if pp is not None:
        for border in list(pp.findall(qn('w:pBdr'))): pp.remove(border)
if 'ReportTable' not in d.styles: d.styles.add_style('ReportTable', WD_STYLE_TYPE.TABLE)
if 'Normal Table' in d.styles: d.styles['ReportTable'].base_style = d.styles['Normal Table']
d.styles['ReportTable'].font.name = 'Calibri'; d.styles['ReportTable'].font.size = Pt(10.5)


def p(text, style=None): return d.add_paragraph(text, style)
def heading(text, page=False, level=1):
    item = p(text, f'Heading {level}'); item.paragraph_format.page_break_before = page
    return item

def formula(label, value):
    item = p(f'{label}: {value}')
    item.alignment = WD_ALIGN_PARAGRAPH.CENTER
    item.paragraph_format.space_before = Pt(3)
    item.paragraph_format.space_after = Pt(10)
    return item

def borders(table):
    pr = table._tbl.tblPr
    old = pr.find(qn('w:tblBorders'))
    if old is not None: pr.remove(old)
    b = OxmlElement('w:tblBorders')
    for side in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); b.append(e)
    pr.append(b)

def table(rows):
    t = d.add_table(rows=0, cols=2); t.style = 'ReportTable'; t.autofit = False
    for col, width in zip(t.columns, [58, 114]): col.width = Mm(width)
    for label, value in rows:
        cells = t.add_row().cells
        for cell, width in zip(cells, [58, 114]):
            cell.width = Mm(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.space_before = Pt(3)
        cells[0].text, cells[1].text = label, value
        cells[0].paragraphs[0].runs[0].bold = True
        t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    borders(t)
    return t

# Restore the reference's branded header and document identification footer.
for part in [s.header, s.footer]:
    for child in list(part._element): part._element.remove(child)
header = s.header.add_paragraph(); header.paragraph_format.space_after = Pt(0)
header.paragraph_format.tab_stops.add_tab_stop(Mm(172), WD_TAB_ALIGNMENT.RIGHT)
header.add_run().add_picture(str(ROOT / 'Template files/report-assets/customer-logo.png'), width=Mm(95))
header.add_run('\t')
header.add_run().add_picture(str(ROOT / 'Template files/report-assets/acs-logo.png'), width=Mm(60))
footer = s.footer.add_table(rows=3, cols=3, width=Mm(172)); footer.autofit = False
footer.alignment = WD_TABLE_ALIGNMENT.CENTER
footer.cell(0, 0).text = 'ISSUED BY  {{PREPAREDBY}}'
footer.cell(0, 1).text = 'PROJECT MANAGER  {{MANAGER}}'
footer.cell(0, 2).text = 'GENERATED  {{REPORTDATE}}'
footer.cell(1, 0).text = 'SHEET  Cable tray calculations {{TRAYNAME}}'
footer.cell(1, 1).merge(footer.cell(1, 2)).text = 'PROJECT  {{PROJECTNAME}}'
footer.cell(2, 0).merge(footer.cell(2, 1)).text = 'PROJECT NUMBER  {{PROJECTNUMBER}}'
page = footer.cell(2, 2).paragraphs[0]; page.add_run('Page ')
for field_name, suffix in [('PAGE', ' of '), ('NUMPAGES', '')]:
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), field_name)
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '1'; run.append(text); field.append(run); page._p.append(field); page.add_run(suffix)
for row in footer.rows:
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(0); para.paragraph_format.line_spacing = 1
            for run in para.runs: run.font.size = Pt(8)
borders(footer)
s.footer.add_paragraph().paragraph_format.space_after = Pt(0)

# Cover follows the reference: identification first, then tray and support dimensions.
cover = p('Cable tray report', 'Title'); cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_before = Pt(26); cover.paragraph_format.space_after = Pt(25)
name = p('Cable tray name  {{TRAYNAME}}', 'Subtitle'); name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.paragraph_format.space_after = Pt(20)
for text in ['Cable tray type  {{TRAYMANUFACTURER}} {{TRAYTYPE}}', 'Cable tray purpose  {{TRAYPURPOSE}}']:
    p(text).alignment = WD_ALIGN_PARAGRAPH.CENTER
heading('Cable tray dimensions')
p('Height: {{TRAYHEIGHT}}, Width: {{TRAYWIDTH}}, Length: {{TRAYLENGTH}}, Own mass per metre: {{TRAYOWNPERMETER}}')
heading('Support dimensions')
p('Support type: {{SUPPORTTYPE}}, Length: {{SUPPORTLENGTH}}, Unit mass: {{SUPPORTWEIGHT}} kg, Selected spacing: {{DISTANCE}}')
p('This report records the routed cable schedule, explains the weight and installation-width calculations, and checks cable-plus-tray loading against the assigned manufacturer curve. Assessment results are collected before the final cable laying concept.')

heading('Cables laying on the tray', True)
p('The following schedule lists the cables routed through this tray, their diameters and their unit masses. These values form the input to the weight calculation. The cable laying concept on the final page shows their bundle arrangement.')
p('{{CABLESTABLE}}')

heading('Weight calculations')
heading('Supports weight calculations', level=2)
p('The number of supports depends on route length L and selected support spacing s. For tray type {{TRAYTYPE}}, the calculation uses s = {{DISTANCE}}. A route shorter than this spacing has two supports. For longer routes, n is the whole number of spans floor(L / s), and the base count is max(2, n + 1). One additional support is included when n is at least one and the remaining length L - n x s exceeds 20% of s.')
formula('Supports count', '{{SUPPORTSCOUNT}}')
p('The total support mass is obtained by multiplying the support count by the mass of one support. This mass is retained in the route weight inventory.')
formula('Supports total mass', '{{SUPPTOTALWEIGHT}}')
p('The equivalent support mass per metre is obtained by dividing the total support mass by the full tray route length.')
formula('Supports mass per metre', '{{SUPPWEIGHTPERMETER}}')

heading('Tray own weight calculations', level=2)
p('The tray own mass per metre is {{TRAYOWNPERMETER}}. For the route weight inventory, the equivalent support mass per metre is added to this value. Support mass is excluded from the separate manufacturer-curve check because it is not carried by the tray span.')
formula('Tray and supports mass per metre', '{{TRAYLOADPERMETER}}')
p('The total mass of the tray and its supports is obtained by multiplying their combined mass per metre by the full route length.')
formula('Tray and supports total mass', '{{TRAYCALCWEIGHT}}')

heading('Cables on tray weight calculations', level=2)
p('The cable mass per metre is the sum of the unit masses of all the routed cables in the schedule. Each cable contributes once; the selected external grounding cable is added once when enabled.')
p('{{GROUNDINGNOTE}}')
formula('Cables mass per metre', '{{CABLEWEIGHTSUM}}')
p('The total cable mass on the reported route is the cable mass per metre multiplied by the route length. Each routed cable is assumed to occupy this full length.')
formula('Cables total mass', '{{CABLESWEIGHTCALCULATIONS}}')

heading('Total weight', level=2)
p('The combined route inventory adds the tray, support and cable masses. It is reported per metre and as a total for the route.')
formula('Total inventory mass per metre', '{{TOTALPERPOINT}}')
formula('Total inventory mass', '{{TOTALCALC}}')
p('Calculations use the stored input precision. Displayed numbers are rounded, so a small difference may occur when recalculating from displayed values.')

heading('Tray loading calculations', True)
p('The uniformly distributed load used for the manufacturer curve consists of cable mass plus the own mass of the tray. The equivalent support mass used in the route inventory above is not included in this load.')
p('Cable-plus-tray mass per metre: {{SPANMASS}}. The selected support spacing is {{CALCULATEDSPAN}} m.')
heading('Conversion from mass to line load', level=2)
p('The combined cable and tray mass per metre is multiplied by gravitational acceleration g = 9.80665 m/s². Dividing by 1000 converts the resulting force per metre from N/m to kN/m.')
formula('Unfactored line load', '{{SPANLOADFORMULA}}')
heading('Project safety allowance', level=2)
p('The project specifies a safety allowance of {{SAFETYFACTOR}}. The design load is the unfactored line load multiplied by 1 plus the allowance divided by 100.')
formula('Design line load', '{{DESIGNLOADFORMULA}}')
heading('Allowable load from the manufacturer curve', level=2)
p('The assigned curve is {{LOADCURVENAME}}, covering spans {{LOADRANGE}}. At a published span point, its load value is used directly. Between two published points, the allowable load is obtained by linear interpolation using the selected span and the surrounding span/load pairs. No capacity is extrapolated outside the documented range.')
formula('Allowable line load', '{{ALLOWABLELOADFORMULA}}')
p('{{DIAGRAMTRAYPIC}}')
p('Picture 1  Manufacturer load curve with selected support spacing and design load', 'Caption')

heading('Free space calculations', True)
p('The available installation width is determined from the tray geometry and the configured cable bundles. The selected tray type is {{TRAYMANUFACTURER}} {{TRAYTYPE}}. Where rung geometry is specified, the rung height is {{RUNGHEIGHT}} and it reduces the usable height inside the tray.')
p('{{TRAYPICTURE}}')
p('Picture 2  Cable tray type overview', 'Caption')
p('The overall tray side-rail height is {{TRAYHEIGHT}}. The useful height is found by subtracting the rung height from the overall height.')
formula('Useful tray height', '{{USEFULTRAY}}')
p('Free space is the unoccupied installation width across the bottom of the tray. It may lie between bundles, between a bundle and a side rail, or between groups laid along opposite rails. The result is a width percentage, rather than a cable area or volume percentage.')

heading('Space occupied by cables')
p('For a single row, occupied width is the sum of cable diameters plus the spacing between adjacent cables. For bundles arranged in multiple rows, the calculation uses the bundle footprints from the actual Tray details layout; it does not add the diameters of every cable as though all cables were laid in one row.')
p('The configured spacing between cables is {{CABLESPACING}}. The project option to count bundle spacing as free width is {{BUNDLESPACINGFREE}}. Cable widths and the spacing counted as occupied are combined as follows.')
formula('Occupied installation width', '{{OCCUPIEDWIDTHFORMULA}}')
p('The resulting occupied width is {{DIAMETERSSUM}}. An external grounding conductor is excluded from this width when it is mounted outside the side rail.')

heading('Cable tray free space')
p('The free-width percentage is obtained by subtracting occupied width from tray width, dividing by tray width and multiplying by 100. Negative available width is reported as zero free width when the layout exceeds the tray width.')
formula('Percentage of free width', '{{FREESPACE}}')
p('The project free-width limits are a minimum of {{MINFREEWIDTH}} and a maximum of {{MAXFREEWIDTH}}. The final comparison with these limits is included in the assessment results below.')

# All assessment metrics and conclusions are immediately before the final picture page.
heading('Assessment results', True)
heading('Load assessment', level=2)
p('Capacity utilization is the design line load divided by the allowable line load at the selected support spacing, multiplied by 100. The remaining design-load capacity is the allowable load minus the design load. A negative reserve indicates overload.')
formula('Load capacity utilization', '{{LOADUTILIZATIONFORMULA}}')
formula('Remaining design-load capacity', '{{LOADRESERVEFORMULA}}')
table([
    ('Design line load', '{{DESIGNLOAD}} kN/m'),
    ('Allowable line load', '{{ALLOWABLELOAD}} kN/m'),
    ('Load capacity utilization', '{{LOADUTILIZATION}}'),
    ('Remaining design-load capacity', '{{LOADRESERVE}}'),
    ('Remaining capacity percentage', '{{LOADRESERVEPERCENT}}'),
    ('Maximum allowable span within curve range', '{{MAXALLOWABLESPAN}}'),
])
p('{{LOADVERIFICATION}}')
p('This conclusion covers the specified distributed tray load and manufacturer curve. It does not establish the capacity of support brackets, fixings or anchors, or verify concentrated loads.')
heading('Free space assessment', level=2)
p('Calculated free width: {{FREEWIDTHRESULT}}. Project limits: minimum {{MINFREEWIDTH}}, maximum {{MAXFREEWIDTH}}.')
p('{{FREEWIDTHVERIFICATION}}')

heading('Cable laying concept', True)
p('{{FILLPICTURE}}')
p('Picture 3  Cable distribution and bundle arrangement along the selected tray', 'Caption')
update_fields = d.settings.element.find(qn('w:updateFields'))
if update_fields is None:
    update_fields = OxmlElement('w:updateFields'); d.settings.element.append(update_fields)
update_fields.set(qn('w:val'), 'true')
d.core_properties.title = 'Cable tray report'
d.core_properties.subject = 'LV tray weight load and free-space calculations'
d.core_properties.author = ''; d.core_properties.last_modified_by = ''
d.save(TARGET)
print(TARGET)
